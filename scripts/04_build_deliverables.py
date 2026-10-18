"""Build the report, saved-output notebook, dashboard and narrated presentation."""
from pathlib import Path
import hashlib
import re
import sys
import json
import shutil
import nbformat as nbf
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'scripts'))
from study_content import sections, slides, evidence
from src.dashboard import build as build_dashboard
from src.exhibits import build as build_figures
from src.annotation_packet import build as build_human_packet
from src.reader_evidence import build as build_reader_evidence
from src.part_b_demography import build as build_demography
from src.rubric_evidence import build as build_rubric_evidence
from src.country_birth_rates import build as build_country_birth_rates


def esc(value):
    mapping = {'&': r'\&', '%': r'\%', '$': r'\$', '#': r'\#', '_': r'\_',
               '{': r'\{', '}': r'\}', '~': r'\textasciitilde{}', '^': r'\textasciicircum{}'}
    return ''.join(mapping.get(char, char) for char in str(value))


def prose(value):
    saved = []
    def keep(text):
        saved.append(text)
        return f'ZZTOKEN{len(saved)-1}ZZ'
    value = re.sub(r'\[([^]]+)\]\((https?://[^)]+)\)',
                   lambda m: keep(r'\href{' + esc(m[2]) + '}{' + esc(m[1]) + '}'), value)
    value = re.sub(r'\*\*(.*?)\*\*', lambda m: keep(r'\textbf{' + esc(m[1]) + '}'), value)
    value = esc(value)
    for i, text in enumerate(saved):
        value = value.replace(f'ZZTOKEN{i}ZZ', text)
    return value


def latex_table(item):
    n = len(item['headers'])
    widths = item.get('widths', [1 / n] * n)
    layout = ''.join(r'>{\raggedright\arraybackslash}p{' + str(w) +
                     r'\dimexpr\linewidth-' + str(2*n) + r'\tabcolsep\relax}' for w in widths)
    lines = [r'\par\smallskip\noindent{\sffamily\fontsize{9.0}{11.1}\selectfont\setlength{\fboxsep}{0pt}%',
             r'\colorbox{mist}{\begin{tabular}{' + layout + '}',
             ' & '.join((r'\rule[-4pt]{0pt}{17pt}' if i == 0 else '') +
                         r'\textbf{\color{navy}' + esc(h) + '}' for i, h in enumerate(item['headers'])) + r' \\',
             r'\end{tabular}}\par\nobreak\nointerlineskip\noindent%',
             r'\begin{tabular}{' + layout + '}', r'\toprule[0.3pt]']
    lines.extend(' & '.join((r'\textbf{' + esc(x) + '}') if i == 0 else esc(x)
                            for i, x in enumerate(row)) + r' \\[3pt]' for row in item['rows'])
    return '\n'.join(lines + [r'\bottomrule[0.4pt]\end{tabular}}\par\smallskip'])


PREAMBLE = r'''\documentclass[11pt,letterpaper]{article}
\usepackage[margin=0.7in,headheight=21pt,headsep=10pt,footskip=20pt]{geometry}
\usepackage{fontspec}
\IfFontExistsTF{Cambria}{\setmainfont{Cambria}[Ligatures=NoCommon]}{\setmainfont{Latin Modern Roman}[Ligatures=NoCommon]}
\IfFontExistsTF{Calibri}{\setsansfont{Calibri}}{\setsansfont{Latin Modern Sans}}
\IfFontExistsTF{Georgia}{\newfontfamily\coverfont{Georgia}}{\newfontfamily\coverfont{Latin Modern Roman}}
\usepackage{xcolor}
\usepackage{amsmath,booktabs,array,graphicx,fancyhdr,titlesec,hyperref}
\definecolor{navy}{HTML}{172F43}\definecolor{teal}{HTML}{167D8D}
\definecolor{slate}{HTML}{657787}\definecolor{mist}{HTML}{EAF1F4}
\definecolor{rulecolor}{HTML}{CDDCE2}\definecolor{coral}{HTML}{C77B68}
\hypersetup{colorlinks=true,urlcolor=teal,linkcolor=navy,pdfauthor={Panagiotis Housos},pdftitle={Economic Narratives and Consumer Sentiment Ahead of the 2026 Midterms}}
\color{navy}
\pagestyle{fancy}\fancyhf{}
\fancyhead[L]{\sffamily\fontsize{8.3}{10}\selectfont\color{slate} AI APPLICATIONS IN FINANCE\quad /\quad RESEARCH REPORT}
\fancyhead[R]{\sffamily\fontsize{8.3}{10}\selectfont\color{slate} Panagiotis Housos\enspace |\enspace ph2606}
\fancyfoot[L]{\sffamily\fontsize{8.0}{10}\selectfont\color{slate} Economy: October 13, 2026\enspace |\enspace Reader extension: October 18, 2026}
\fancyfoot[R]{\sffamily\small\color{teal}\bfseries\ifnum\value{page}<10 0\fi\thepage\enspace\textcolor{slate}{/ TOTALPAGES}}
\renewcommand{\headrulewidth}{0.35pt}
\renewcommand{\headrule}{\hbox to\headwidth{\color{rulecolor}\leaders\hrule height \headrulewidth\hfill}}
\titleformat{\section}{\sffamily\fontsize{18}{21}\selectfont\bfseries\color{navy}}{}{0pt}{}
\titlespacing*{\section}{0pt}{0pt}{10pt}
\setlength{\parindent}{0pt}\setlength{\parskip}{7pt}
\setlength{\tabcolsep}{4pt}\renewcommand{\arraystretch}{1.05}
\emergencystretch=2em
\begin{document}
\fontsize{10.6}{13.5}\selectfont
'''


def build_report(content):
    tex = [PREAMBLE.replace('TOTALPAGES', str(len(content)))]
    data = evidence()
    latest = data['primary'].iloc[-1]
    cards = [(f"{int(data['media'].item_count.sum()):,}", 'Headlines analyzed'),
             (str(int(data['comparison'].n.iloc[0])), 'Shared forecast test months'),
             (f'{latest.gap:.1f}', 'September: actual - prediction')]
    figure_number = 0
    markdown = ['# Economic Narratives and Consumer Sentiment Ahead of the 2026 Midterms',
                'Panagiotis Housos | ph2606 | Fall 2026', '']
    for page, (title, items) in enumerate(content):
        if page:
            tex.append(r'\newpage')
        tex.append(r'\fontsize{10.6}{13.5}\selectfont\setlength{\parskip}{6pt}')
        if page == 0:
            tex += [r'{\sffamily\fontsize{9}{11}\selectfont\bfseries\color{teal} ECONOMIC PERCEPTIONS\quad /\quad 2026 MIDTERMS}\par\vspace{5pt}',
                    r'{\coverfont\fontsize{30}{33}\selectfont\color{navy} Economic narratives\\[-1pt]\& consumer sentiment}\par\vspace{5pt}',
                    r'{\sffamily\fontsize{11}{14}\selectfont\color{slate} Ahead of the 2026 U.S. midterms\enspace /\enspace With a manifesto for independent judgment}\par\vspace{8pt}',
                    r'\begingroup\setlength{\fboxsep}{8pt}\noindent']
            for i, (value, label) in enumerate(cards):
                tex.append(r'\colorbox{mist}{\begin{minipage}[c][43pt][c]{.29\linewidth}\raggedright'
                           + r'{\sffamily\fontsize{24}{27}\selectfont\bfseries\color{navy} ' + esc(value) + r'}\par\vspace{2pt}'
                           + r'{\sffamily\fontsize{8.2}{10}\selectfont\color{slate} ' + esc(label)
                           + r'}\end{minipage}}' + (r'\hfill' if i < 2 else ''))
            tex.append(r'\par\endgroup\vspace{12pt}')
        numbered = re.match(r'^(\d+)\.\s+(.*)', title)
        heading = (r'\textcolor{teal}{\fontsize{11}{13}\selectfont ' + numbered[1].zfill(2)
                   + r'}\quad ' + esc(numbered[2])) if numbered else esc(title)
        tex.append(r'\section*{' + heading + '}')
        markdown += ['## ' + title, '']
        for item in items:
            if item['type'] == 'p':
                tex += [prose(item['text']), r'\par']
                markdown += [item['text'], '']
            elif item['type'] == 'table':
                tex.append(latex_table(item))
                markdown += ['| ' + ' | '.join(item['headers']) + ' |',
                             '| ' + ' | '.join(['---'] * len(item['headers'])) + ' |']
                markdown += ['| ' + ' | '.join(map(str, row)) + ' |' for row in item['rows']]
                markdown.append('')
            elif item['type'] == 'figure':
                figure_number += 1
                options = 'width=' + item.get('width', '.97') + r'\linewidth'
                if item.get('height'):
                    options += ',height=' + item['height'] + ',keepaspectratio'
                tex += [r'\begin{center}\includegraphics[' + options + ']{outputs/' + item['name'] + r'.pdf}\end{center}',
                        r'\noindent\begingroup\setlength{\fboxsep}{5pt}\colorbox{mist}{\parbox{\dimexpr\linewidth-10pt\relax}{'
                        + r'\sffamily\fontsize{8.5}{10.8}\selectfont\color{navy}\textbf{FIGURE '
                        + f'{figure_number:02d}' + r'}\quad ' + prose(item['caption']) + r'}}\endgroup\par\smallskip']
                markdown += ['![' + item['caption'] + '](outputs/' + item['name'] + '.png)', '']
    (ROOT / 'report.tex').write_text('\n'.join(tex + [r'\end{document}']), encoding='utf-8')
    (ROOT / 'REPORT.md').write_text('\n'.join(markdown), encoding='utf-8')
    (ROOT/'outputs/report_structure.json').write_text(json.dumps({
        'sections':[title for title,_ in content], 'expected_pages':len(content),
        'page_limit_override':'User explicitly requested an integrated Part B and instructed disregarding the 10-12 page limit.',
        'part_b_start_page':next(i+1 for i,(title,_) in enumerate(content) if title.startswith('Part B')),
        'appendices_start_page':next(i+1 for i,(title,_) in enumerate(content) if title.startswith('Appendix A')),
    },indent=2),encoding='utf-8')


def build_notebook(content):
    input_paths = ['data/processed/monthly_dataset.csv','data/processed/media_monthly.csv',
                   'outputs/tables/provider_bias.csv','data/processed/investment_prices.csv',
                   'data/processed/media_scores.csv','data/processed/bias_trading_prices.csv',
                   'data/processed/part_b_demography.csv','data/processed/country_birth_rates_2023.csv']
    expected_inputs = {name:hashlib.sha256((ROOT/name).read_bytes()).hexdigest()
                       for name in input_paths if (ROOT/name).exists()}
    cells = [nbf.v4.new_markdown_cell('# Economic narratives and consumer sentiment\n\n**Panagiotis Housos | ph2606 | AI Applications in Finance | Final Project**\n\nInformation cutoff: October 13, 2026. Human validation is not completed; independent AI reviews are explicitly separate. Acquisition and model scoring run separately. The cells below rerun analysis from local frozen inputs. Raw datasets are excluded from Git under the established submission policy.'),
             nbf.v4.new_code_cell("from pathlib import Path\nimport json\nimport hashlib\nimport pandas as pd\nfrom IPython.display import display, Image\nROOT = Path.cwd()\nassert (ROOT / 'src/economics.py').exists(), 'Execute from the project root.'\nexpected_inputs = " + repr(expected_inputs) + "\nfor name, expected in expected_inputs.items():\n    path = ROOT / name\n    assert path.exists(), f'Acquire required input: {name}'\n    assert hashlib.sha256(path.read_bytes()).hexdigest() == expected, f'Input changed: {name}. Rerun analysis and rebuild deliverables before executing this notebook.'\npd.set_option('display.max_columns', 14)\npd.set_option('display.max_colwidth', 70)\npd.set_option('display.precision', 4)"),
             nbf.v4.new_markdown_cell('## Recompute the frozen empirical analysis\n\nThe economic feature matrix uses historical ALFRED snapshots. Media reconstruction is retrospective and subject to source selection and original-wording limitations. The package never invents missing records, exposure or human labels.'),
             nbf.v4.new_code_cell("import subprocess, sys\nfor command in [['scripts/03_analyze.py'], ['-m', 'src.dynamics'], ['-m', 'src.investing', 'analyze'], ['-m', 'src.bias_trading', 'analyze'], ['scripts/08_validation_tables.py'], ['-m','src.external_validation','score'], ['-m','src.rubric_evidence'], ['-m','src.part_b_demography','build'], ['-m','src.reader_evidence'], ['-m','src.country_birth_rates','build']]:\n    result = subprocess.run([sys.executable, *command], cwd=ROOT, capture_output=True, text=True)\n    print(result.stdout)\n    assert result.returncode == 0, result.stderr\nprint('Primary study and separately identified extensions recalculated from frozen inputs.')")]
    cells += [nbf.v4.new_markdown_cell('## Persistence and differenced-model sensitivity\n\nThese checks were added after the primary analysis and are exploratory. ADF and KPSS have different null hypotheses; structural breaks limit either interpretation. The differenced models use the same 44 test months and a separately identified two-contrast loss-test family.'),
              nbf.v4.new_code_cell("display(pd.read_csv(ROOT / 'outputs/tables/dynamics_unit_root.csv'))\ndisplay(pd.read_csv(ROOT / 'outputs/tables/dynamics_differenced_metrics.csv'))\ndisplay(pd.read_csv(ROOT / 'outputs/tables/dynamics_loss_hac.csv'))")]
    cells += [nbf.v4.new_markdown_cell('## Independent AI diagnostic: class and source concordance\n\nThese are AI development references, not human ground truth. The original pool has 300 items; final-corpus eligibility and mixed/uncertain exclusions are explicit. Small source/class cells have unstable estimates. Required human validation remains pending.'),
              nbf.v4.new_code_cell("ai = json.loads((ROOT/'outputs/ai_diagnostic_validation_metrics.json').read_text())\ndisplay(pd.Series({k: ai.get(k) for k in ['label_type','completed_independent_pairs','tone_agreement','tone_cohens_kappa','annotation_pool_final_eligible_n','annotation_pool_final_excluded_n','human_specific_requirement_satisfied']}))\nfor name in ['cohorts','model_concordance','class_concordance','source_concordance','confusion_matrices','dimension_agreement']:\n    print('Explicit original and final populations:', name)\n    display(pd.read_csv(ROOT / ('outputs/tables/validation_'+name+'.csv')))"),
              nbf.v4.new_markdown_cell('## Political-name and source-label stress tests\n\nThe original text is held constant except for the specified change. Changed names can change meaning. Class switches do not prove political discrimination or factual bias.'),
              nbf.v4.new_code_cell("media_summary = json.loads((ROOT / 'outputs/media_summary.json').read_text())\ndisplay(pd.DataFrame(media_summary['perturbation']['variant_comparisons']))\nprint('Source-prefix paired mean absolute tone difference:', media_summary['perturbation']['source_pair_mean_absolute_delta'])\nprint('Factual-calibration status:', media_summary.get('factual_calibration'))")]
    cells += [nbf.v4.new_markdown_cell('## External human-consensus transfer benchmark\n\nFinancial PhraseBank company/investor sentences differ from US economic titles. The frozen class-balanced 300 examples and unknown training overlap prevent claims of a pristine holdout. These are published human labels, not two new annotators for this project. Target-domain human validation remains incomplete.'),
              nbf.v4.new_code_cell("for name in ['models','classes','confusion']:\n    display(pd.read_csv(ROOT/('outputs/tables/external_validation_'+name+'.csv')))\nexternal = json.loads((ROOT/'outputs/tables/external_validation_summary.json').read_text())\nprint('Limitations:', *external['limitations'], sep='\\n- ')\nassert external['human_specific_requirement_satisfied'] is False"),
              nbf.v4.new_markdown_cell('## Statistical evidence and prospective power\n\nFull-history coefficient associations, chronological forecast comparisons and sample-size planning are different estimands. The planned experiment has not been run.'),
              nbf.v4.new_code_cell("for name in ['loss_contrasts','tone_associations','survey_components','economic_sensitivities','media_sensitivities','dynamic_sensitivities','forecast_skill','power_plan']:\n    print(name)\n    display(pd.read_csv(ROOT/('outputs/tables/rubric_'+name+'.csv')))"),
              nbf.v4.new_markdown_cell('## Part B source data and published summaries\n\nDemography measures births, not cognition. The AI learning chart transcribes published regression summaries; it is not a new participant-level analysis. The manifesto is an argument distinct from Part A estimates.'),
              nbf.v4.new_code_cell("display(pd.read_csv(ROOT/'outputs/tables/part_b_demography_comparison.csv'))\nwith pd.option_context('display.max_rows', 100):\n    display(pd.read_csv(ROOT/'outputs/tables/country_birth_rates_top100.csv'))\ndisplay(pd.read_csv(ROOT/'outputs/tables/reader_ai_published_effects.csv'))\ndisplay(pd.read_csv(ROOT/'outputs/tables/reader_attention_published.csv'))\ndisplay(pd.read_csv(ROOT/'outputs/tables/reader_evidence_sources.csv'))")]
    if (ROOT/'outputs/tables/provider_bias.csv').exists():
        cells += [nbf.v4.new_markdown_cell('## News-provider comparison: matched official releases\n\nA separate supplemental sample includes CNBC and Fox brands. The provider bars use a common set of releases and a defined peer reference. This is a relative negative-framing proxy, not a factual or ideological rating. The original forecast corpus and 300-title AI review remain separate; the new providers do not have independent human validation.'),
                  nbf.v4.new_code_cell("provider_table = pd.read_csv(ROOT / 'outputs/tables/provider_bias.csv')\ndisplay(provider_table)\nprovider_summary = json.loads((ROOT / 'outputs/tables/provider_bias_summary.json').read_text())\ndisplay(pd.Series(provider_summary, name='provider_comparison_protocol'))")]
    cells += [nbf.v4.new_markdown_cell('## Retrospective investment extension (October 17)\n\nThe primary macro/trend rule was fixed before performance calculation; the simpler monthly 60% SPY / 40% BIL benchmark is preferred after the comparison. Neither is an untouched holdout. Costs include entry, exit and both traded legs. Gap and media overlays use retrospective inputs. All 44 gaps are negative; this is exposure, not demonstrated timing.'),
              nbf.v4.new_code_cell("investment = json.loads((ROOT / 'outputs/tables/investment_summary.json').read_text())\nperformance = pd.read_csv(ROOT / 'outputs/tables/investment_metrics.csv')\ndisplay(performance.loc[performance.cost_bps.eq(10)])\nprint('Full-period cost sensitivity:')\ndisplay(performance.loc[performance['sample'].eq('full'), ['model','cost_bps','net_cagr','max_drawdown','annual_gross_turnover']])\nprint('Primary HAC diagnostics:')\ndisplay(pd.DataFrame(investment['hac_tests']).T)\nprint('Original two-test overlay family (unadjusted for exposure):')\ndisplay(pd.DataFrame(investment['overlay_tests']))\nprint('Post-result exposure diagnosis:')\nprint(json.dumps(investment['exposure_diagnostics'], indent=2))\nprint(investment['practical_conclusion'])")]
    cells += [nbf.v4.new_code_cell("display(Image(filename=str(ROOT / 'outputs/investment_strategy.png')))"),
              nbf.v4.new_markdown_cell('## Three aggressive strategies that use only earlier news wording\n\nA strategy is a rule for choosing what to buy. These three ideas use the same national-minus-local negative-wording score. None uses economic releases, survey answers or stock-price trends to choose funds. UPRO aims for three times the S&P 500\'s DAILY move; SPXU aims for three times the opposite DAILY move. These are risky daily targets, not promises for an entire month. Their actual historical fund prices are used. The fixed rules, current-month exclusion, costs and all benchmark results appear below. A losing idea is retained. These historical simulations do not establish verified bias or a future investment advantage.'),
              nbf.v4.new_code_cell("bias = json.loads((ROOT / 'outputs/tables/bias_trading_summary.json').read_text())\nbias_metrics = pd.read_csv(ROOT / 'outputs/tables/bias_trading_metrics.csv')\nprint('All rules and all benchmarks, on exactly the same dates:')\ndisplay(bias_metrics.loc[bias_metrics.cost_bps.eq(10)])\nprint('Cost sensitivity for the three ideas:')\ndisplay(bias_metrics.loc[bias_metrics['sample'].eq('full') & bias_metrics.model.isin(bias['rule_names']), ['model','cost_bps','net_cagr','net_total_return','max_drawdown','terminal_value_10000']])\nprint('How often each rule chose each fund:')\ndisplay(pd.DataFrame(bias['rule_fund_counts']).fillna(0).astype(int))\nprint('Three historical return comparisons, with Holm correction (not alpha tests):')\ndisplay(pd.DataFrame(bias['tests_vs_spy']))\nprint('Signal inputs only:', bias['signal_inputs_only'])\nprint('Limitations:', *bias['limitations'], sep='\\n- ')")]
    for page_index,(title, items) in enumerate(content):
        cells.append(nbf.v4.new_markdown_cell('## ' + title))
        for item_index,item in enumerate(items):
            if item['type'] == 'p':
                cells.append(nbf.v4.new_markdown_cell(item['text']))
            elif item['type'] == 'table':
                code = ("import sys\nif str(ROOT / 'scripts') not in sys.path: sys.path.insert(0, str(ROOT / 'scripts'))\n"
                        "from study_content import sections\nitem = sections()["+str(page_index)+"][1]["+str(item_index)+"]\n"
                        "display(pd.DataFrame(item['rows'], columns=item['headers']))")
                cells.append(nbf.v4.new_code_cell(code))
            elif item['type'] == 'figure':
                cells.append(nbf.v4.new_code_cell("display(Image(filename=str(ROOT / 'outputs' / " + repr(item['name'] + '.png') + ")))"))
                cells.append(nbf.v4.new_markdown_cell(item['caption']))
    cells.append(nbf.v4.new_markdown_cell('## Reproducibility and limitations\n\nSee `README.md`, `SNAPSHOT.md`, `AI_USE.md`, the source manifest, codebook, causal proposal, and annotation-validation protocol. A fresh download can change mutable feed results. Saved aggregate outputs preserve this snapshot, while item-level data remain local.'))
    nb = nbf.v4.new_notebook(cells=cells)
    nb.metadata.kernelspec = {'name': 'finance_final_project', 'display_name': 'Python (Finance final project)', 'language': 'python'}
    nb.metadata.language_info = {'name': 'python', 'version': '3.12'}
    nbf.validate(nb)
    nbf.write(nb, ROOT / 'final_project.ipynb')


def add_text(slide, left, top, width, height, text, size=24, color='172F43', bold=False, paragraph_space=14):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_top = Pt(0)
    frame.margin_bottom = Pt(0)
    lines = text.split('\n')
    for i, line in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = 'Calibri'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(paragraph_space if i < len(lines)-1 else 0)
    return box


def add_picture_fit(slide, path, left, top, width, height):
    picture = slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))
    if picture.height > Inches(height):
        picture.width = round(picture.width * Inches(height) / picture.height)
        picture.height = Inches(height)
    picture.left = Inches(left) + (Inches(width)-picture.width)//2
    picture.top = Inches(top) + (Inches(height)-picture.height)//2
    return picture


def build_slides(content):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    notes = []
    for i, item in enumerate(content):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string('FCFBF8')
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(.055))
        accent.fill.solid()
        accent.fill.fore_color.rgb = RGBColor.from_string('167D8D')
        accent.line.fill.background()
        add_text(slide, .55, .35, 12.2, .9, item['title'], 32, bold=True)
        if item.get('layout') == 'strategy_rules':
            add_text(slide,.65,1.18,12.0,.70,item['intro'],19,paragraph_space=0)
            add_text(slide,.65,2.00,12.0,.70,item['funds'],17,paragraph_space=0)
            for column,(label,title,body) in enumerate(item['cards']):
                left=.65+column*4.12
                card=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(left),Inches(2.88),Inches(3.78),Inches(2.22))
                card.fill.solid();card.fill.fore_color.rgb=RGBColor.from_string('EAF1F4')
                card.line.fill.background()
                add_text(slide,left+.18,3.03,3.4,.28,label,12,color='167D8D',bold=True)
                add_text(slide,left+.18,3.42,3.4,.42,title,20,bold=True)
                add_text(slide,left+.18,4.02,3.4,1.0,body,18,paragraph_space=0)
            add_text(slide,.68,5.38,11.95,1.0,item['example'],18,paragraph_space=0)
            add_text(slide,.68,6.55,11.95,.26,item['caption'],11,color='657787')
        elif item.get('layout') == 'wide_chart':
            add_text(slide, .60, 1.18, 12.1, .40, item['subtitle'], 20)
            add_picture_fit(slide, ROOT / 'outputs' / (item['image'] + '.png'), .65, 1.73, 12.0, 4.65)
            add_text(slide, .65, 6.55, 12.0, .35, item['caption'], 13)
        else:
            add_text(slide, .58, 1.38, 5.5 if item.get('image') else 11.9, 5.1,
                     '\n'.join(item['points']), 23)
        if item.get('image') and item.get('layout') != 'wide_chart':
            add_picture_fit(slide, ROOT / 'outputs' / (item['image'] + '.png'), 6.2, 1.45, 6.6, 4.9)
        add_text(slide, .6, 6.98, 11.7, .3,
                 f'Panagiotis Housos | ph2606   •   Economy cutoff: October 13, 2026   •   {i+1}/{len(content)}', 11, color='657787')
        slide.notes_slide.notes_text_frame.text = item['notes']
        notes += [f"Slide {i+1}: {item['title']}", item['notes'], '']
    prs.save(ROOT / 'Panagiotis_Housos_Final_Project.pptx')
    (ROOT / 'docs' / 'presentation_script.md').write_text('\n'.join(notes), encoding='utf-8')
    word_count = sum(len(s['notes'].split()) for s in content)
    print(f'Presentation: {len(content)} slides, {word_count} narration words (~{word_count/140:.1f} min at 140 wpm).')


if __name__ == '__main__':
    (ROOT / 'docs').mkdir(exist_ok=True)
    build_human_packet()
    build_figures()
    build_reader_evidence()
    build_demography()
    build_rubric_evidence()
    build_country_birth_rates()
    content = sections()
    assert len(content) > 12 and any(title.startswith('Part B') for title,_ in content)
    build_report(content)
    build_notebook(content)
    build_dashboard()
    build_slides(slides())
    if (ROOT/'outputs/provider_bias.png').exists():
        shutil.copyfile(ROOT/'outputs/provider_bias.png',ROOT/'Provider_Bias_Comparison.png')
        shutil.copyfile(ROOT/'outputs/provider_bias.pdf',ROOT/'Provider_Bias_Comparison.pdf')
    shutil.copyfile(ROOT/'outputs/investment_strategy.png',ROOT/'Investment_Strategy_Comparison.png')
    shutil.copyfile(ROOT/'outputs/investment_strategy.pdf',ROOT/'Investment_Strategy_Comparison.pdf')
    shutil.copyfile(ROOT/'outputs/bias_trading.png',ROOT/'Bias_Trading_Comparison.png')
    shutil.copyfile(ROOT/'outputs/bias_trading.pdf',ROOT/'Bias_Trading_Comparison.pdf')
    for source, target in [('part_b_fertility_trends','Reader_Fertility_Trends'),
                           ('reader_ai_learning','Reader_AI_Learning')]:
        for extension in ['png','pdf']:
            shutil.copyfile(ROOT/'outputs'/f'{source}.{extension}',ROOT/f'{target}.{extension}')
    for chart in range(1,6):
        shutil.copyfile(ROOT/'outputs'/f'country_birth_rates_{chart:02d}.png',
                        ROOT/f'Country_Birth_Rates_{chart:02d}.png')
    shutil.copyfile(ROOT/'outputs/Country_Birth_Rates_Top_100.pdf',ROOT/'Country_Birth_Rates_Top_100.pdf')
    print(f'Built {len(content)}-section integrated report source, notebook, dashboard and presentation.')
