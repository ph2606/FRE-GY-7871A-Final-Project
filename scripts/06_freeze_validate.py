"""Validate deliverable consistency and bind the submission to local input hashes."""
from pathlib import Path
from datetime import datetime, timezone
import argparse
import base64
import hashlib
import importlib.metadata as md
import json
import platform
import re
import subprocess
import sys
import nbformat
import pandas as pd
from packaging.requirements import Requirement
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs'


def sha(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024*1024), b''):
            h.update(chunk)
    return h.hexdigest()


def environment_lock():
    required = []
    for line in (ROOT / 'requirements.txt').read_text().splitlines():
        if line.strip() and not line.startswith('#'):
            required.append(Requirement(line).name)
    seen, versions = set(), {}
    while required:
        name = required.pop()
        key = name.lower().replace('_','-')
        if key in seen:
            continue
        seen.add(key)
        dist = md.distribution(name)
        versions[dist.metadata['Name']] = dist.version
        for raw in dist.requires or []:
            requirement = Requirement(raw)
            if requirement.marker is None or requirement.marker.evaluate({'extra': ''}):
                required.append(requirement.name)
    content = ['# Exact observed Python 3.12 Windows environment dependency closure.',
               '# For Linux/macOS regenerate platform-specific packages from requirements.txt.',
               '--extra-index-url https://download.pytorch.org/whl/cpu']
    content += [name + '==' + version for name,version in sorted(versions.items(),key=lambda x:x[0].lower())]
    (ROOT / 'requirements-lock.txt').write_text('\n'.join(content)+'\n',encoding='utf-8',newline='\n')
    return versions


def run_tests():
    result = subprocess.run([sys.executable, '-m', 'pytest', 'tests', '-q'], cwd=ROOT,
                            capture_output=True, text=True)
    print(result.stdout, end='')
    if result.stderr:
        print(result.stderr, end='')
    passed = re.search(r'(\d+) passed', result.stdout)
    if result.returncode != 0 or passed is None:
        raise RuntimeError('Test run failed; submission validation stopped')
    OUT.mkdir(exist_ok=True)
    (OUT/'test_summary.json').write_text(json.dumps({
        'command':'python -m pytest tests -q', 'exit_code':result.returncode,
        'passed':int(passed.group(1)), 'recorded_at_utc':datetime.now(timezone.utc).isoformat()
    }, indent=2), encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run-tests', action='store_true',
                        help='Run and record tests before freezing; required on a fresh reproduction.')
    args = parser.parse_args()
    if args.run_tests:
        run_tests()
    OUT.mkdir(exist_ok=True)
    pdf = ROOT / 'Panagiotis_Housos_Final_Project.pdf'
    reader = PdfReader(pdf)
    pages = len(reader.pages)
    structure=json.loads((OUT/'report_structure.json').read_text())
    assert pages == structure['expected_pages'] and pages > 12, (pages,structure)
    assert structure['page_limit_override'].startswith('User explicitly requested')
    report_text = '\n'.join(page.extract_text() for page in reader.pages)
    assert 'anchoring' in report_text.lower(), 'Required mechanism discussion missing'
    assert 'October 15, 2026' in report_text, 'Supplement acquisition date missing'
    assert 'Part B' in report_text and 'independent judgment' in report_text
    assert '53.7%' in report_text and '0.949' in report_text, 'Validation/statistical evidence missing'
    tests = json.loads((OUT/'test_summary.json').read_text(encoding='utf-8'))
    assert tests['exit_code'] == 0 and tests['passed'] > 0, 'Tests have not passed'
    nb = nbformat.read(ROOT / 'final_project.ipynb', as_version=4)
    nbformat.validate(nb)
    codes = [c for c in nb.cells if c.cell_type == 'code']
    assert codes and all(c.execution_count is not None for c in codes), 'Notebook not fully executed'
    assert not [o for c in codes for o in c.get('outputs',[]) if o.output_type == 'error']
    prs = Presentation(ROOT / 'Panagiotis_Housos_Final_Project.pptx')
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 10
            assert shape.top + shape.height <= prs.slide_height + 10
    words = sum(len(s.notes_slide.notes_text_frame.text.split()) for s in prs.slides)
    assert 6 <= words/140 <= 8, f'Narration duration {words/140:.2f} minutes'
    native_path=OUT/'presentation_native_validation.json'
    native=json.loads(native_path.read_text()) if native_path.exists() else None
    if native:
        assert native['pptx_sha256']==sha(ROOT/'Panagiotis_Housos_Final_Project.pptx')
        assert native['text_overflows']==[] and native['slides']==len(prs.slides)
    html = (ROOT / 'dashboard.html').read_text(encoding='utf-8')
    browser_path=OUT/'browser_validation.json'
    browser=json.loads(browser_path.read_text()) if browser_path.exists() else None
    if browser:
        assert browser['html_sha256']==sha(ROOT/'dashboard.html')
        assert browser['runtime_exceptions']==[] and len(browser['charts'])==10
    assert 'Plotly.newPlot' in html and 'Human validation outstanding' in html
    assert 'Extra negative titles' in html and 'CNBC' in html and 'Fox News' in html
    chart_ids = re.findall(r'<div id="(chart-\d+)"', html)
    assert len(chart_ids) == len(set(chart_ids)) == 10, 'Missing or duplicated dashboard charts'
    provider=pd.read_csv(OUT/'tables/provider_bias.csv')
    if 'comparison_status' in provider:
        provider=provider.loc[provider.comparison_status.eq('main_common_event_panel')]
    assert {'CNBC','Fox News'}.issubset(set(provider.provider))
    assert provider.n_events.nunique()==1 and provider.n_events.min()>=6
    assert provider.n_titles.gt(0).all()
    assert (provider.lower_95<=provider.upper_95).all()
    assert abs(provider.bias_pp.sum())<1e-7, 'Common-release peer contrasts must center at zero'
    assert (ROOT/'Provider_Bias_Comparison.png').exists()
    provider_image = (ROOT/'Provider_Bias_Comparison.png').read_bytes()
    assert provider_image == (OUT/'provider_bias.png').read_bytes()
    assert any(shape.image.blob == provider_image for slide in prs.slides for shape in slide.shapes
               if shape.shape_type == MSO_SHAPE_TYPE.PICTURE), 'Stale provider chart in slides'
    figure_cells = [c for c in codes if 'provider_bias.png' in c.source]
    assert len(figure_cells) == 1
    assert any(base64.b64decode(o.get('data',{}).get('image/png','')) == provider_image
               for o in figure_cells[0].get('outputs',[])), 'Stale provider chart in notebook'
    acquisition = json.loads((OUT/'tables/provider_bias_acquisition.json').read_text(encoding='utf-8'))
    assert acquisition['input_titles_sha256'] == sha(ROOT/'data/processed/provider_bias_titles.csv')
    investment=json.loads((OUT/'tables/investment_summary.json').read_text(encoding='utf-8'))
    assert investment['primary_allocation_periods']==141 and investment['common_allocation_periods']==44
    assert investment['cost_bps_per_dollar_traded']==10 and investment['entry_and_final_liquidation_charged']
    assert investment['protocol_sha256']==sha(ROOT/'docs/INVESTMENT_PROTOCOL.md')
    assert all(x['protocol_sha256_before_performance']==investment['protocol_sha256'] for x in investment['market_retrieval'])
    assert all(x['end']=='2026-10-02' for x in investment['full_metrics'])
    assert investment['exposure_diagnostics']['gap_negative_origins']==44
    for name,digest_value in investment['hashes'].items():
        assert sha(ROOT/name)==digest_value, 'Stale investment input: '+name
    signals=pd.read_csv(ROOT/'data/processed/investment_signals.csv',parse_dates=['month','execution_date','trend_reference_month'])
    assert (signals.execution_date>signals.month).all() and (signals.trend_reference_month<signals.month).all()
    assert 'SPY' in report_text and 'BIL' in report_text
    strategy_image=(ROOT/'Investment_Strategy_Comparison.png').read_bytes()
    assert strategy_image==(OUT/'investment_strategy.png').read_bytes()
    strategy_cells=[c for c in codes if 'investment_strategy.png' in c.source]
    assert len(strategy_cells)==1 and any(base64.b64decode(o.get('data',{}).get('image/png',''))==strategy_image
               for o in strategy_cells[0].get('outputs',[])), 'Stale investment chart in notebook'
    bias=json.loads((OUT/'tables/bias_trading_summary.json').read_text(encoding='utf-8'))
    assert bias['allocation_periods']==68 and bias['cost_bps']==10 and bias['entry_and_liquidation_included']
    assert bias['protocol_sha256']==sha(ROOT/'docs/BIAS_TRADING_PROTOCOL.md')
    assert all(row['protocol_sha256_before_performance']==bias['protocol_sha256']
               for row in bias['acquisition']['sources'])
    assert bias['script_sha256']==sha(ROOT/'src/bias_trading.py')
    assert bias['signal_script_sha256']==sha(ROOT/'src/bias_signal.py')
    for name,digest_value in bias['hashes'].items():assert sha(ROOT/name)==digest_value,name
    assert all(row['start']=='2021-02-02' and row['end']=='2026-10-02' for row in bias['full_metrics'])
    assert len(bias['tests_vs_spy'])==3 and all('holm_pvalue' in row for row in bias['tests_vs_spy'])
    bias_signals=pd.read_csv(ROOT/'data/processed/bias_trading_signals.csv',parse_dates=['month','signal_month','execution_date'])
    assert (bias_signals.signal_month<bias_signals.month).all() and (bias_signals.execution_date>bias_signals.month).all()
    assert bias_signals.baseline_n.eq(12).all()
    assert all(set(bias_signals[model])<={'UPRO','SPXU','BIL'} for model in bias['rule_names'])
    assert all(term in report_text for term in ['UPRO','SPXU','Fade','Follow','Wait'])
    bias_image=(ROOT/'Bias_Trading_Comparison.png').read_bytes()
    assert bias_image==(OUT/'bias_trading.png').read_bytes()
    assert any(shape.image.blob==bias_image for slide in prs.slides for shape in slide.shapes
               if shape.shape_type==MSO_SHAPE_TYPE.PICTURE), 'Stale news-only chart in slides'
    bias_cells=[c for c in codes if 'bias_trading.png' in c.source]
    assert len(bias_cells)==1 and any(base64.b64decode(o.get('data',{}).get('image/png',''))==bias_image
               for o in bias_cells[0].get('outputs',[])), 'Stale news-only chart in notebook'
    panel=pd.read_csv(ROOT/'data/processed/monthly_dataset.csv')
    assert panel.loc[panel.month>'2026-09-01','sentiment'].isna().all(), 'Final target published after cutoff'
    metrics = pd.read_csv(OUT/'tables/model_metrics.csv')
    common = metrics[(metrics['sample']=='common_media_oos')&(metrics.target=='sentiment')]
    assert common.n.nunique()==1 and common.start.nunique()==1 and common.end.nunique()==1
    assert {'M1_economic','M2_economic_media','M3_economic_media_controls','Persistence'}.issubset(set(common.model))
    forecasts = pd.read_csv(OUT/'tables/forecasts.csv',parse_dates=['month','training_end'])
    assert (forecasts.training_end < forecasts.month).all()
    media = pd.read_csv(ROOT/'data/processed/media_monthly.csv')
    assert media.month.nunique() == len(media) and media.media_tone.between(-1,1).all()
    human = json.loads((OUT/'human_validation_metrics.json').read_text())
    assert human['human_specific_requirement_satisfied'] is False
    assert human['completed_independent_pairs'] == 0
    external=json.loads((OUT/'tables/external_validation_summary.json').read_text())
    assert external['human_specific_requirement_satisfied'] is False
    assert all(row['n']==300 for row in external['comparisons'].values())
    assert external['protocol_sha256']==sha(ROOT/'docs/EXTERNAL_VALIDATION_PROTOCOL.md')
    assert external['freeze_sha256']==sha(ROOT/'data/processed/external_phrasebank_freeze.json')
    assert external['predictions_sha256']==sha(ROOT/'data/processed/external_phrasebank_predictions.csv')
    demographic=json.loads((OUT/'tables/part_b_demography_summary.json').read_text())
    demography=pd.read_csv(ROOT/'data/processed/part_b_demography.csv')
    assert len(demography)==18*34 and demographic['last_year']==2023
    for source,target in [('part_b_demography_comparison','Reader_Demography_Comparison'),
                          ('part_b_fertility_trends','Reader_Fertility_Trends'),
                          ('reader_ai_learning','Reader_AI_Learning')]:
        blob=(ROOT/f'{target}.png').read_bytes()
        assert blob==(OUT/f'{source}.png').read_bytes()
        assert (ROOT/f'{target}.pdf').read_bytes()==(OUT/f'{source}.pdf').read_bytes()
        if source!='part_b_fertility_trends':
            assert any(shape.image.blob==blob for slide in prs.slides for shape in slide.shapes
                       if shape.shape_type==MSO_SHAPE_TYPE.PICTURE), 'Stale Part B slide chart'
        cells=[c for c in codes if f'{source}.png' in c.source]
        assert len(cells)==1 and any(base64.b64decode(o.get('data',{}).get('image/png',''))==blob
                   for o in cells[0].get('outputs',[])), 'Stale Part B notebook chart'
    versions = environment_lock()
    inventory = []
    for folder in ['data','outputs']:
        for path in sorted((ROOT/folder).rglob('*')):
            if not path.is_file() or any(x in path.parts for x in ['__pycache__','jupyter_runtime','ipython','edge_preview']):
                continue
            if path.name in ['snapshot_manifest.json','deliverable_validation.json']:
                continue
            inventory.append({'path':path.relative_to(ROOT).as_posix(),'bytes':path.stat().st_size,'sha256':sha(path)})
    digest = hashlib.sha256(json.dumps(inventory,sort_keys=True,separators=(',',':')).encode()).hexdigest()
    (OUT/'snapshot_manifest.json').write_text(json.dumps(inventory,indent=2),encoding='utf-8')
    summary = json.loads((OUT/'tables/economic_summary.json').read_text())
    media_summary = json.loads((OUT/'media_summary.json').read_text())
    lines = ['# Submission snapshot','',
             'Information cutoff: **October 13, 2026**, America/New_York. Election day is outside this snapshot.',
             'Raw/derived datasets, detailed annotations and model weights remain local and are excluded from Git.',
             'A fresh download reproduces the method, but mutable search feeds, original wording, deletions and revisions can change results.','',
             f'- Local file inventory: {len(inventory):,} files, {sum(x["bytes"] for x in inventory):,} bytes.',
             f'- Canonical inventory SHA-256: `{digest}`.',
             '- Full itemized inventory is local at `outputs/snapshot_manifest.json`.',
             f'- Media: {media_summary["n_titles"]:,} titles, {media_summary["months"]} months, {media_summary["complete_months"]} complete.',
             f'- Common media assessment: {summary["media_comparison"]["holdout_n"]} months.',
             '- Human-specific annotation requirement: **not satisfied**; AI diagnostics are separately identified.',
             '- October 18 integrated revision: Part B manifesto, primary-source evidence and World Bank fertility/birth-rate exhibits; report page cap explicitly overridden by user.',
             '- External published human-consensus benchmark: 300 company-news sentences, separate from target-corpus annotation; Cardiff accuracy53.7%, VADER51.0%, with important domain and training-overlap limits.',
             f'- Provider supplement: {len(provider)} providers and {int(provider.n_events.iloc[0])} identical releases; collected retrospectively for the October 15 revision.',
             '- Investing extension: developed retrospectively October 17; 141 monthly allocations, January 2, 2015–October 2, 2026; 10 bp per dollar traded including entry and liquidation.',
             '- The 60% SPY / 40% BIL benchmark is preferred after the comparison; positive historical profit is not demonstrated media alpha or a promise of future profit.',
             '- Final plain-language revision adds three news-framing-only strategies using observed +3x/-3x daily funds, 68 monthly decisions and all six same-date benchmarks; no guaranteed bias advantage is claimed.',
             '', '| Artifact / key input | SHA-256 |', '| --- | --- |']
    keys = ['data/processed/monthly_dataset.csv','data/processed/media_corpus.csv','data/processed/media_scores.csv',
            'data/processed/media_monthly.csv','outputs/tables/forecasts.csv','outputs/tables/model_metrics.csv',
            'outputs/tables/predictive_loss_tests.csv','docs/MEDIA_CODEBOOK.md',
            'outputs/tables/provider_bias.csv','outputs/tables/provider_bias_summary.json','Provider_Bias_Comparison.png','Provider_Bias_Comparison.pdf',
            'data/processed/investment_prices.csv','data/processed/investment_signals.csv','outputs/tables/investment_metrics.csv',
            'outputs/tables/investment_summary.json','docs/INVESTMENT_PROTOCOL.md','Investment_Strategy_Comparison.png','Investment_Strategy_Comparison.pdf',
            'data/processed/bias_trading_prices.csv','data/processed/bias_trading_monthly.csv','data/processed/bias_trading_signals.csv',
            'outputs/tables/bias_trading_metrics.csv','outputs/tables/bias_trading_summary.json','docs/BIAS_TRADING_PROTOCOL.md',
            'Bias_Trading_Comparison.png','Bias_Trading_Comparison.pdf',
            'data/processed/part_b_demography.csv','outputs/tables/part_b_demography_summary.json',
            'outputs/tables/external_validation_summary.json','docs/EXTERNAL_VALIDATION_PROTOCOL.md',
            'outputs/tables/validation_summary.json','outputs/tables/rubric_evidence_summary.json',
            'outputs/tables/reader_evidence_summary.json','outputs/report_structure.json',
            'Reader_Demography_Comparison.png','Reader_Demography_Comparison.pdf',
            'Reader_Fertility_Trends.png','Reader_Fertility_Trends.pdf','Reader_AI_Learning.png','Reader_AI_Learning.pdf',
            'Panagiotis_Housos_Final_Project.pdf','Panagiotis_Housos_Final_Project.pptx','final_project.ipynb','dashboard.html']
    for name in keys:
        path = ROOT/name
        if path.exists():lines.append(f'| `{name}` | `{sha(path)}` |')
    lines += ['',f'Observed interpreter: Python {platform.python_version()} on {platform.system()}; exact package closure in requirements-lock.txt.',
              'The latest Michigan historical targets are not historical survey vintages. Macro predictors use actual ALFRED snapshots; media tests remain retrospective pseudo-out-of-sample.',
              'Detailed AI annotations and claim checks are development diagnostics, not validated human ground truth.']
    (ROOT/'SNAPSHOT.md').write_text('\n'.join(lines)+'\n',encoding='utf-8',newline='\n')
    result = {'validated_at_utc':datetime.now(timezone.utc).isoformat(),'report_pages':pages,
              'notebook_code_cells':len(codes),'notebook_errors':0,'presentation_slides':len(prs.slides),
              'narration_words':words,'estimated_minutes_140wpm':words/140,
              'common_forecast_months':int(common.n.iloc[0]),'snapshot_manifest_sha256':digest,
              'provider_comparison_n':len(provider),'common_provider_events':int(provider.n_events.iloc[0]),
              'tests_passed':tests['passed'],'dashboard_charts':len(chart_ids),
              'investment_allocation_months':141,'investment_return_end':'2026-10-02',
              'investment_cost_bps':10,'preferred_allocation':'60% SPY / 40% BIL; retrospective preference',
              'bias_trading_months':68,'bias_trading_rules':3,'bias_trading_return_end':'2026-10-02',
              'human_requirement_satisfied':False,'platform':platform.platform(),
              'part_b_integrated':True,'report_page_limit_override':structure['page_limit_override'],
              'native_powerpoint_review_bound':bool(native),'browser_review_bound':bool(browser),
              'python':platform.python_version(),'package_count':len(versions)}
    (OUT/'deliverable_validation.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result,indent=2))


if __name__=='__main__':main()
