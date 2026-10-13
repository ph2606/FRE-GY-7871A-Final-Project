"""Download official data and actual historical ALFRED monthly vintages.

No API credentials required. Existing raw files are immutable cache entries.
Run: python scripts/01_collect_economics.py [--workers 8]
"""
from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys
import time

import openpyxl
import pandas as pd
import requests

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.economics import ROOT, CUTOFF, SERIES, OCTOBER_PRELIMINARY, parse_alfred, parse_michigan, sha256, vintage_features, write_json

RAW = ROOT / "data/raw/economics"
OUT = ROOT / "data/processed"
TABLES = ROOT / "outputs/tables"


def acquire(url: str, path: Path, validator=None) -> dict:
    path.parent.mkdir(parents=True, exist_ok=True)
    metadata_path = path.with_suffix(path.suffix + ".meta.json")
    if path.exists() and metadata_path.exists():
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        if metadata.get('sha256') != sha256(path):
            raise ValueError(f'Cached economics payload hash mismatch: {path.name}')
        if metadata.get('url') != url:
            raise ValueError(f'Cached economics request URL mismatch: {path.name}')
        if validator:
            validator(path.read_text(encoding="utf-8-sig"))
        return metadata
    failure = None
    for attempt in range(4):
        try:
            response = requests.get(url, timeout=(15, 60), headers={"User-Agent": "Academic-Finance-Project/1.0 (public-data-research)"})
            response.raise_for_status()
            if validator:
                validator(response.text)
            path.write_bytes(response.content)
            result = {"url": url, "path": str(path.relative_to(ROOT)), "retrieved_at_utc": datetime.now(timezone.utc).isoformat(),
                      "status": response.status_code, "sha256": sha256(path), "bytes": len(response.content),
                      "content_type": response.headers.get("Content-Type"), "last_modified": response.headers.get("Last-Modified")}
            write_json(metadata_path, result)
            return result
        except Exception as exc:
            failure = str(exc)
            time.sleep(min(1 + attempt, 3))
    raise RuntimeError(f"Acquisition failed {url}: {failure}")


def vintage_job(series_id: str, date: str) -> dict:
    # 48 months are sufficient for the largest lag (36-month price change),
    # including a buffer for delayed BEA releases during shutdowns.
    begin = (pd.Timestamp(date) - pd.DateOffset(months=48)).strftime("%Y-%m-%d")
    url = f"https://alfred.stlouisfed.org/graph/alfredgraph.csv?id={series_id}&cosd={begin}&coed={date}&vintage_date={date}"
    return acquire(url, RAW / "vintages" / series_id / f"{date}.csv", lambda text: parse_alfred(text, series_id, date))


def audit_workbook(official: dict[str, pd.Series], core: pd.Series) -> dict:
    path = ROOT.parent / "Perceptual_Bias_Data.xlsx"
    if not path.exists():
        result = {"status": "optional_source_workbook_unavailable", "workbook": path.name,
                  "treatment": "Official panel estimation proceeds; reproducing the supplied-file audit requires the original workbook in the parent directory."}
        write_json(TABLES / "workbook_audit.json", result)
        return result
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    records = []
    mappings = [(0, 1, "unemployment"), (3, 4, "sentiment"), (7, 8, "pce_yoy"), (10, 11, "gdp_qoq_annualized")]
    overview = []
    for date_col, value_col, key in mappings:
        values = [(r[date_col], r[value_col]) for r in list(sheet.values)[1:] if r[date_col] is not None and r[value_col] is not None]
        dates = [pd.Timestamp(x[0]).to_period("M").to_timestamp() for x in values]
        overview.append({"measure": key, "n": len(values), "min": str(min(dates).date()), "max": str(max(dates).date()),
                         "duplicate_dates": len(dates) - len(set(dates)), "original_excel_date_is_observation_period_end": True})
        for original, value in values:
            date = pd.Timestamp(original).to_period("M").to_timestamp()
            if key == "gdp_qoq_annualized":
                date = date.to_period("Q").start_time
            official_value = official[key].get(date)
            status = "matched" if official_value is not None and pd.notna(official_value) and abs(float(value) - float(official_value)) <= (.11 if key in ("unemployment", "sentiment", "gdp_qoq_annualized") else .02) else "mismatch_or_revision"
            if date > official[key].index.max():
                status = "unverified_period_after_latest_official_observation"
            if key == "sentiment" and date == pd.Timestamp("2026-10-01"):
                official_value = OCTOBER_PRELIMINARY["sentiment"]
                status = "matches_verified_preliminary_not_final"
            if key == "gdp_qoq_annualized" and date >= pd.Timestamp("2026-07-01"):
                status = "not_released_by_project_cutoff_exclude"
            records.append({"measure": key, "excel_period_end": original, "observation_period": date,
                            "workbook_value": value, "official_cutoff_vintage_value": official_value,
                            "difference": float(value) - float(official_value) if official_value is not None and pd.notna(official_value) else None,
                            "core_pce_yoy_crosscheck": core.get(date) if key == "pce_yoy" else None,
                            "status": status})
    frame = pd.DataFrame(records)
    frame.to_csv(TABLES / "workbook_audit.csv", index=False)
    recent = frame.loc[frame.observation_period >= pd.Timestamp("2026-01-01")].copy()
    recent["official_cutoff_vintage_value"] = recent["official_cutoff_vintage_value"].astype(object).where(recent["official_cutoff_vintage_value"].notna(), None)
    recent["difference"] = recent["difference"].astype(object).where(recent["difference"].notna(), None)
    recent["core_pce_yoy_crosscheck"] = recent["core_pce_yoy_crosscheck"].astype(object).where(recent["core_pce_yoy_crosscheck"].notna(), None)
    pce = frame.loc[frame.measure == "pce_yoy"].dropna(subset=["core_pce_yoy_crosscheck"])
    core_matches = int((abs(pce.workbook_value - pce.core_pce_yoy_crosscheck) < .0001).sum())
    result = {"workbook": path.name, "sha256": sha256(path), "sheet": sheet.title, "dimensions": [sheet.max_row, sheet.max_column],
              "separate_date_columns": True, "publication_dates_present": False, "vintage_metadata_present": False,
              "treatment": "User workbook audited as supplied; official sources replace it for estimation. Rows joined by observation date, never by row.",
              "series": overview, "status_counts": frame.groupby(["measure", "status"]).size().rename("n").reset_index().to_dict("records"),
              "pce_definition_audit": {"compared_rows": len(pce), "matches_core_within_0_0001_percentage_points": core_matches,
                  "conclusion": "Workbook PCE YoY tracks core PCE (PCEPILFE), whereas assignment requires headline PCE (PCEPI); use recomputed headline series."},
              "recent_rows": recent.to_dict("records"),
              "specific_flags": ["Workbook Q3 2026 GDP 3.7 predates the October 29 advance release; exclude.",
                                 "Workbook October sentiment 46.3 matches the official October 9 PRELIMINARY release; keep separate from final-only model evaluation.",
                                 "Workbook PCE YoY matches core PCE instead of required headline PCE; recalculate headline inflation from PCEPI.",
                                 "End-of-month Excel dates identify observation periods and cannot be treated as publication dates.",
                                 "2025 October BLS U-3 is missing in official data because the household survey was not collected; do not silently fill."]}
    write_json(TABLES / "workbook_audit.json", result)
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workers", type=int, default=8)
    parser.add_argument("--latest-only", action="store_true")
    args = parser.parse_args()
    for path in (RAW, OUT, TABLES):
        path.mkdir(parents=True, exist_ok=True)
    origins = pd.date_range("2010-01-01", "2026-10-01", freq="MS")
    tasks = [(s, d.strftime("%Y-%m-%d")) for d in origins for s in SERIES] if not args.latest_only else []
    tasks += [(s, CUTOFF.strftime("%Y-%m-%d")) for s in SERIES]
    manifest, failures = [], []
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = {pool.submit(vintage_job, series, date): (series, date) for series, date in tasks}
        for n, future in enumerate(as_completed(futures), 1):
            try:
                manifest.append(future.result())
            except Exception as exc:
                failures.append({"series": futures[future][0], "date": futures[future][1], "error": str(exc)})
            if n % 80 == 0 or n == len(tasks):
                print(f"ALFRED snapshots {n}/{len(tasks)}, failures={len(failures)}", flush=True)
    write_json(TABLES / "economic_acquisition_failures.json", failures)
    if failures:
        raise RuntimeError(f"{len(failures)} vintage downloads failed; rerun to use cache and retry.")
    official_urls = {
        "michigan_headline.csv": "https://www.sca.isr.umich.edu/files/tbmics.csv",
        "michigan_components.csv": "https://www.sca.isr.umich.edu/files/tbmiccice.csv",
        "bea_pce_latest.html": "https://www.bea.gov/news/2026/personal-income-and-outlays-august-2026",
        "bea_gdp_latest.html": "https://www.bea.gov/data/gdp/gross-domestic-product",
        "bea_release_schedule.html": "https://www.bea.gov/news/schedule",
        "bls_employment_september2026.html": "https://www.bls.gov/news.release/archives/empsit_10022026.htm",
        "bls_u3_2009_2018.json": "https://api.bls.gov/publicAPI/v2/timeseries/data/LNS14000000?startyear=2009&endyear=2018",
        "bls_u3_2019_2026.json": "https://api.bls.gov/publicAPI/v2/timeseries/data/LNS14000000?startyear=2019&endyear=2026",
    }
    for filename, url in official_urls.items():
        try:
            manifest.append(acquire(url, RAW / filename))
        except Exception as exc:
            failures.append({"series": filename, "date": str(CUTOFF.date()), "error": str(exc)})
    headline = parse_michigan(RAW / "michigan_headline.csv").rename(columns={"ICS_ALL": "sentiment"})
    components = parse_michigan(RAW / "michigan_components.csv")
    # Official component header order is ICC (current) then ICE (expectations).
    components = components.rename(columns={components.columns[0]: "current_conditions", components.columns[1]: "expectations"})
    michigan = headline.join(components)
    michigan = michigan.loc[michigan.index <= CUTOFF].copy()
    michigan.to_csv(OUT / "michigan_official.csv")
    if not args.latest_only:
        rows = []
        for origin in origins:
            date = origin.strftime("%Y-%m-%d")
            values = {s: parse_alfred((RAW / "vintages" / s / f"{date}.csv").read_text(encoding="utf-8"), s, date) for s in SERIES}
            rows.append(vintage_features(values, date))
        frame = pd.DataFrame(rows)
        frame["month"] = pd.to_datetime(frame["month"])
        frame = frame.merge(michigan.reset_index(), on="month", how="left")
        for target in ["sentiment", "current_conditions", "expectations"]:
            # Monthly Michigan final tables; target history has not been vintaged.
            # Explicitly separate this limitation from vintaged macro predictors.
            frame[f"{target}_lag1"] = frame["month"].map(michigan[target].rename(index=lambda x: x + pd.DateOffset(months=1)))
        frame["month_sin"] = __import__("numpy").sin(2 * __import__("numpy").pi * frame["month"].dt.month / 12)
        frame["month_cos"] = __import__("numpy").cos(2 * __import__("numpy").pi * frame["month"].dt.month / 12)
        frame["post_web_transition"] = (frame["month"] >= "2024-07-01").astype(int)
        frame.to_csv(OUT / "monthly_dataset.csv", index=False)
    date = CUTOFF.strftime("%Y-%m-%d")
    latest_values = {s: parse_alfred((RAW / "vintages" / s / f"{date}.csv").read_text(encoding="utf-8"), s, date) for s in SERIES}
    latest = vintage_features(latest_values, date)
    latest["sentiment_latest_period"] = str(michigan.index.max().date())
    latest.update(michigan.iloc[-1].astype(float).to_dict())
    latest["october_preliminary"] = OCTOBER_PRELIMINARY
    write_json(TABLES / "economic_latest.json", latest)
    write_json(RAW / "michigan_october2026_preliminary_web_transcription.json", OCTOBER_PRELIMINARY)
    write_json(TABLES / "michigan_october_preliminary.json", OCTOBER_PRELIMINARY)
    # Audit all workbook history using the same project-cutoff vintage, not a
    # four-year feature slice.
    full = {}
    for series in SERIES:
        url = f"https://alfred.stlouisfed.org/graph/alfredgraph.csv?id={series}&cosd=1977-01-01&coed={date}&vintage_date={date}"
        file = RAW / f"{series}_cutoff_full.csv"
        manifest.append(acquire(url, file, lambda text, s=series: parse_alfred(text, s, date)))
        full[series] = parse_alfred(file.read_text(encoding="utf-8"), series, date)
    official = {"unemployment": full["UNRATE"], "sentiment": michigan["sentiment"],
                "pce_yoy": 100 * (full["PCEPI"] / full["PCEPI"].shift(12) - 1),
                "gdp_qoq_annualized": 100 * ((full["GDPC1"] / full["GDPC1"].shift(1)) ** 4 - 1)}
    core_url = f"https://alfred.stlouisfed.org/graph/alfredgraph.csv?id=PCEPILFE&cosd=1977-01-01&coed={date}&vintage_date={date}"
    core_file = RAW / "PCEPILFE_cutoff_full.csv"
    manifest.append(acquire(core_url, core_file, lambda text: parse_alfred(text, "PCEPILFE", date)))
    core = parse_alfred(core_file.read_text(encoding="utf-8"), "PCEPILFE", date)
    core_yoy = 100 * (core / core.shift(12) - 1)
    audit_workbook(official, core_yoy)
    # Cross-check cutoff U-3 against data obtained directly from the BLS API.
    bls_records = []
    for filename in ("bls_u3_2009_2018.json", "bls_u3_2019_2026.json"):
        payload = json.loads((RAW / filename).read_text(encoding="utf-8"))
        if payload.get("status") != "REQUEST_SUCCEEDED":
            raise ValueError(f"BLS request did not succeed: {filename}")
        for row in payload["Results"]["series"][0]["data"]:
            if row["period"].startswith("M") and row["period"] != "M13":
                period = pd.Timestamp(year=int(row["year"]), month=int(row["period"][1:]), day=1)
                value = pd.to_numeric(row["value"], errors="coerce")
                bls_records.append({"period": period, "bls_u3": value, "alfred_cutoff_u3": full["UNRATE"].get(period), "footnotes": json.dumps(row.get("footnotes", []))})
    pd.DataFrame(bls_records).to_csv(TABLES / "bls_alfred_crosscheck.csv", index=False)
    write_json(TABLES / "economic_acquisition_failures.json", failures)
    # A --latest-only run must preserve the full historical cache manifest.
    manifest = [json.loads(p.read_text(encoding="utf-8")) for p in RAW.rglob("*.meta.json")]
    transcription = RAW / "michigan_october2026_preliminary_web_transcription.json"
    manifest.append({"url": OCTOBER_PRELIMINARY["source_url"], "path": str(transcription.relative_to(ROOT)),
                     "retrieved_at_utc": None, "verified_on_client_date": "2026-10-13",
                     "sha256": sha256(transcription), "bytes": transcription.stat().st_size,
                     "acquisition_method": OCTOBER_PRELIMINARY["acquisition_method"]})
    pd.DataFrame(manifest).to_csv(TABLES / "economic_provenance.csv", index=False)
    write_json(TABLES / "economic_data_dictionary.json", {"cutoff": date, "series": SERIES,
        "forecast_origin": "first calendar day of target month; actual macro vintages at that date",
        "quarterly_alignment": "latest GDP quarter actually present in vintage; no interpolation",
        "pce_yoy_formula": "100 * (PCEPI[t] / PCEPI[t-12 calendar months] - 1)",
        "gdp_qoq_annualized_formula": "100 * ((GDPC1[q] / GDPC1[q-1])**4 - 1)",
        "gdp_yoy_formula": "100 * (GDPC1[q] / GDPC1[q-4] - 1)",
        "historical_publication_dates": "Not individually recovered. Actual availability is observed in monthly vintage snapshots; publication is bounded above by snapshot date.",
        "target_vintage_limitation": "Michigan final headline/components from latest official tables, without historical target vintages; preliminaries excluded.",
        "latest_publication_dates": [{"measure": "U-3 September 2026", "published": "2026-10-02 08:30 America/New_York", "source": official_urls["bls_employment_september2026.html"]},
                                     {"measure": "PCE August 2026; real income August 2026", "published": "2026-09-30 08:30 America/New_York", "source": official_urls["bea_pce_latest.html"]},
                                     {"measure": "GDP Q2 2026 third estimate", "published": "2026-09-30 08:30 America/New_York", "source": official_urls["bea_gdp_latest.html"]},
                                     {"measure": "GDP Q3 2026 advance (future)", "published": "2026-10-29 08:30 America/New_York", "source": official_urls["bea_release_schedule.html"]}]})
    print(json.dumps({"latest": latest, "provenance_files": len(manifest), "optional_failures": failures}, indent=2), flush=True)


if __name__ == "__main__":
    main()
